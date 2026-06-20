# cv_parser.py
import os
import pdfplumber
import docx
import instructor
from groq import Groq
from pydantic import BaseModel, Field
from typing import List, Optional

# ---------------------- Structural Sub-Schemas ----------------------

class ExperienceSchema(BaseModel):
    company: str = Field(description="Name of the organization, company, or workplace")
    role: str = Field(description="Job title, role, or designation held")
    duration: str = Field(description="Time period or duration of the job (e.g., Oct 2025 - Present)")
    achievements: List[str] = Field(description="Core responsibilities, tasks performed, or major metrics achieved")

class EducationSchema(BaseModel):
    institution: str = Field(description="Name of the school, college, or university")
    degree: str = Field(description="Exact degree name or certification (e.g., BSCS, ICS, Matriculation)")
    passing_year: Optional[str] = Field(None, description="Graduation year or completion date")
    gpa: Optional[str] = Field(None, description="CGPA, grade, or percentage obtained")

class ProjectSchema(BaseModel):
    title: str = Field(description="Title or name of the project")
    tech_stack: List[str] = Field(description="List of tools, programming languages, or technologies used")
    description: str = Field(description="A brief summary detailing features or goals of the project")

# ---------------------- Main Composite Schema ----------------------

class CVAnalysisSchema(BaseModel):
    full_name: str = Field(description="Full legal or professional name of the candidate")
    email: str = Field(description="Primary contact email address")
    phone: Optional[str] = Field(None, description="Contact phone or mobile number if present")
    links: List[str] = Field(description="Social profile links such as GitHub URL, LinkedIn profile, or portfolio sites")
    skills: List[str] = Field(description="Comprehensive list of tools, frameworks, technical skills, and soft skills mentioned")
    education: List[EducationSchema] = Field(description="Structured list of all educational milestones")
    experience: List[ExperienceSchema] = Field(description="Structured list of all professional work experiences")
    projects: List[ProjectSchema] = Field(description="Structured list of all personal or group projects showcased")
    experience_years: float = Field(default=0.0, description="Calculated total years of professional work experience represented as a float (e.g., 2.5 or 5.0).")

# ---------------------- Extraction Utility Functions ----------------------

def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Safely extracts raw plain text string from a PDF file using pdfplumber layout engine.
    """
    full_text = ""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text(layout=False) # layout=False ensures a logical text streaming structure
                if text:
                    full_text += text + "\n"
        return full_text.strip()
    except Exception as e:
        print(f"Extraction Error inside cv_parser module: {e}")
        return ""

def extract_text_from_docx(docx_path: str) -> str:
    """
    Safely extracts raw plain text from a DOCX file, including paragraphs and tables.
    """
    try:
        doc = docx.Document(docx_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text).strip()
    except Exception as e:
        print(f"Extraction Error inside cv_parser module (docx): {e}")
        return ""

def extract_text(file_path: str) -> str:
    """
    Unified text extractor for PDF and DOCX files.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    return ""

def parse_cv(file_path: str) -> Optional[CVAnalysisSchema]:
    """
    Transforms unstructured resume text into a strictly validated Pydantic JSON structure using Groq LLM.
    """
    raw_text = extract_text(file_path)
    if not raw_text:
        print("Error: No text extracted from the CV.")
        return None
        
    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        print("Error: GROQ_API_KEY environment variable is missing.")
        return None

    # Initialize Instructor patch over Groq client
    client = instructor.from_groq(Groq(api_key=api_key))
    
    try:
        structured_data = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            response_model=CVAnalysisSchema,
            messages=[
                {
                    "role": "system", 
                    "content": "You are a professional ATS system backend component. Your task is to process raw text extracts from CVs and seamlessly structure every piece of information into the designated target JSON model. Do not truncate, omit, or overlook any project, link, or experience details."
                },
                {
                    "role": "user", 
                    "content": f"Analyze and parse the following resume plain text extract:\n\n{raw_text}"
                }
            ],
            temperature=0.1
        )
        return structured_data
    except Exception as e:
        print(f"Instructor cognitive parsing failed on primary model: {e}. Trying fallback model llama-3.1-8b-instant...")
        try:
            structured_data = client.chat.completions.create(
                model="llama-3.1-8b-instant",
                response_model=CVAnalysisSchema,
                messages=[
                    {
                        "role": "system", 
                        "content": "You are a professional ATS system backend component. Your task is to process raw text extracts from CVs and seamlessly structure every piece of information into the designated target JSON model. Do not truncate, omit, or overlook any project, link, or experience details."
                    },
                    {
                        "role": "user", 
                        "content": f"Analyze and parse the following resume plain text extract:\n\n{raw_text}"
                    }
                ],
                temperature=0.1
            )
            return structured_data
        except Exception as e2:
            print(f"Instructor cognitive parsing failed on fallback model: {e2}")
            return None